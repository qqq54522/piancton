from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
SOURCE = BASE_DIR / "cloud_server_purchase_plan.md"
OUTPUT = BASE_DIR / "Piancton云服务器资源采购方案.docx"

FONT_EAST_ASIA = "Heiti SC"
FONT_LATIN = "Heiti SC"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 38, 46)
MUTED = RGBColor(99, 108, 119)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"


def set_run_font(run, size: float | None = None, bold: bool | None = None, color=None) -> None:
    run.font.name = FONT_LATIN
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360, indent_dxa: int = 120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.1) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_para(doc: Document, text: str = "", style: str | None = None, *, size=11, bold=False, color=INK):
    para = doc.add_paragraph(style=style)
    set_paragraph_spacing(para)
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return para


def add_heading(doc: Document, text: str, level: int) -> None:
    para = doc.add_paragraph()
    if level == 1:
        set_paragraph_spacing(para, before=16, after=8, line=1.1)
        size, color = 16, BLUE
    elif level == 2:
        set_paragraph_spacing(para, before=12, after=6, line=1.1)
        size, color = 13, BLUE
    else:
        set_paragraph_spacing(para, before=8, after=4, line=1.1)
        size, color = 12, DARK_BLUE
    run = para.add_run(text)
    set_run_font(run, size=size, bold=True, color=color)
    para.style = f"Heading {min(level, 3)}"


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=4, line=1.15)
    r = p.add_run(title)
    set_run_font(r, size=11.5, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0, line=1.15)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=INK)
    doc.add_paragraph()


def add_metadata(doc: Document) -> None:
    rows = [
        ("项目", "Piancton 图片素材管理与智能搜索系统"),
        ("用途", "服务器资源采购、预算审批、云厂商选型、上线前资源确认"),
        ("推荐结论", "阿里云 4 核 8GB、100GB 数据盘、按流量公网、峰值 20Mbps"),
        ("版本日期", "V1.0 / 2026-08-31"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=2)
        r1 = p.add_run(f"{label}：")
        set_run_font(r1, size=10.5, bold=True, color=INK)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5, color=INK)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        parts = [part.strip() for part in line.strip().strip("|").split("|")]
        if all(part.replace("-", "").replace(" ", "") == "" for part in parts):
            continue
        rows.append(parts)
    return rows


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            set_cell_margins(cell, top=55, bottom=55, start=110, end=110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                set_cell_shading(cell, LIGHT_FILL)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(paragraph, after=0, line=1.05)
            run = paragraph.add_run(text)
            set_run_font(run, size=8.8, bold=(i == 0), color=INK)
    doc.add_paragraph()


def is_table_line(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def render_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    body_started = False
    table_buffer: list[str] = []

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            add_table(doc, parse_table(table_buffer))
            table_buffer = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("# "):
            continue
        if line.startswith("文档版本") or line.startswith("编制日期") or line.startswith("适用项目") or line.startswith("文档用途"):
            continue
        if line.startswith("## 一、"):
            body_started = True
        if not body_started:
            continue
        if is_table_line(line):
            table_buffer.append(line)
            continue
        flush_table()
        if not line.strip():
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:].strip(), 1)
        elif line.startswith("### "):
            add_heading(doc, line[4:].strip(), 2)
        elif line[:3].count(".") == 1 and line.split(".", 1)[0].isdigit():
            text = line.split(".", 1)[1].strip()
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph_spacing(p, after=5, line=1.15)
            r = p.add_run(text)
            set_run_font(r, size=10.5, color=INK)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph_spacing(p, after=5, line=1.15)
            r = p.add_run(line[2:].strip())
            set_run_font(r, size=10.5, color=INK)
        elif line.startswith("优点：") or line.startswith("缺点：") or line.startswith("优势：") or line.startswith("注意事项：") or line.startswith("综合判断："):
            p = add_para(doc, "", size=10.7, color=INK)
            label, value = line.split("：", 1)
            r1 = p.add_run(f"{label}：")
            set_run_font(r1, size=10.7, bold=True, color=DARK_BLUE)
            r2 = p.add_run(value)
            set_run_font(r2, size=10.7, color=INK)
        else:
            add_para(doc, line.strip(), size=10.7, color=INK)
    flush_table()


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK

    for idx, size, color in ((1, 16, BLUE), (2, 13, BLUE), (3, 12, DARK_BLUE)):
        style = styles[f"Heading {idx}"]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(16 if idx == 1 else 12)
        style.paragraph_format.space_after = Pt(8 if idx == 1 else 6)

    header = section.header.paragraphs[0]
    header.text = "Piancton 云服务器资源采购方案"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(header, after=0)
    for run in header.runs:
        set_run_font(run, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.text = "内部采购与上线准备文件"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(footer, after=0)
    for run in footer.runs:
        set_run_font(run, size=9, color=MUTED)


def build_doc() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    set_paragraph_spacing(title, before=20, after=4)
    r = title.add_run("Piancton 图片素材管理与智能搜索系统")
    set_run_font(r, size=22, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, after=16)
    r = subtitle.add_run("云服务器资源采购方案")
    set_run_font(r, size=15, bold=True, color=DARK_BLUE)

    add_metadata(doc)
    add_callout(
        doc,
        "采购结论",
        "建议首期选择阿里云 4 核 8GB 通用型云服务器，配置 40-80GB 系统盘、100GB 数据盘、按流量公网 IP、20Mbps 峰值带宽，并同步开通对象存储、快照、监控告警、域名和 HTTPS。首期不采购 GPU、不上 Kubernetes、不拆多台服务器。",
    )

    render_markdown(doc, markdown)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
